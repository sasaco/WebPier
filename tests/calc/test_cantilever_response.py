"""The 最大応答震度 read from Aggre-SE's summary workbook (decisions 58, 59): the largest Khmax of both blocks over
every Case, found by the headings, and the workbooks that cannot be read."""

from __future__ import annotations

from pathlib import Path
import re

import pytest

from cantilever_support import workbook
from webpier.calc.cantilever.response import HEADING, SHEET, ResponseError, ResponseValue, read_max_response


def test_the_largest_khmax_of_both_blocks_and_every_case():
    found = read_max_response(workbook({1: ("0.547", "0.552"), 2: ("0.565", "0.571")}), "線路直角方向_Aggre-SE.xlsx")
    assert found.file == "線路直角方向_Aggre-SE.xlsx"
    assert found.values == (
        ResponseValue(1, "線路方向(Case1_ρm10_αf1_正向)", "復旧性", found.values[0].kh),
        ResponseValue(2, "線路方向(Case2_ρm10_αf1_正向)", "復旧性", found.values[1].kh),
        ResponseValue(1, "線路方向(Case1_ρm10_αf1_正向)", "安全性", found.values[2].kh),
        ResponseValue(2, "線路方向(Case2_ρm10_αf1_正向)", "安全性", found.values[3].kh))
    assert [value.kh.value for value in found.values] == [0.547, 0.565, 0.552, 0.571]
    assert found.adopted == found.values[3] and found.kh.value == 0.571 and found.kh.digits == 3


def test_one_case_and_a_tie_keeps_the_earlier_value():
    found = read_max_response(workbook({1: ("0.385", "0.386")}), "線路方向_Aggre-SE.xlsx")
    assert (found.adopted.case, found.adopted.block, found.kh.value) == (1, "安全性", 0.386)
    tie = read_max_response(workbook({1: ("0.400", "0.400"), 2: ("0.400", "0.300")}), "x.xlsx")
    assert (tie.adopted.case, tie.adopted.block) == (1, "復旧性")


def test_the_values_are_read_at_three_decimals():
    found = read_max_response(workbook({1: ("0.3855", "0.38649")}), "x.xlsx")
    assert [value.kh.value for value in found.values] == [0.386, 0.386]
    assert found.adopted.block == "復旧性"


def test_a_workbook_that_is_not_a_zip():
    with pytest.raises(ResponseError, match=re.escape("x.xlsx が xlsx（zip）ではありません")):
        read_max_response(b"not a zip", "x.xlsx")


def test_a_zip_that_is_not_a_workbook():
    import io
    import zipfile
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("fake.txt", "FAKE XLSX")
    with pytest.raises(ResponseError, match=re.escape("x.xlsx を読めません")):
        read_max_response(stream.getvalue(), "x.xlsx")


def test_a_workbook_without_the_sheet():
    with pytest.raises(ResponseError, match=re.escape(f"x.xlsx にシート「{SHEET}」がありません")):
        read_max_response(workbook({1: ("0.3", "0.4")}, sheet="設計総括表"), "x.xlsx")


def test_a_sheet_without_the_case_rows():
    with pytest.raises(ResponseError, match=re.escape(f"x.xlsx の「{SHEET}」に Case の行（No・名称）がありません")):
        read_max_response(workbook({1: ("0.3", "0.4")}, with_cases=False), "x.xlsx")


def test_a_sheet_without_a_case():
    with pytest.raises(ResponseError, match=re.escape(f"x.xlsx の「{SHEET}」に Case がありません")):
        read_max_response(workbook({}), "x.xlsx")


def test_a_sheet_without_the_block_headings():
    with pytest.raises(ResponseError, match=re.escape(f"x.xlsx の「{SHEET}」に 復旧性・安全性 の「{HEADING}」の行がありません")):
        read_max_response(workbook({1: ("0.3", "0.4")}, with_blocks=False), "x.xlsx")


def test_a_value_that_is_not_a_number():
    with pytest.raises(ResponseError, match=re.escape("x.xlsx Case2 安全性 の最大応答震度 '-' を数値として読めません")):
        read_max_response(workbook({1: ("0.3", "0.4"), 2: ("0.3", "-")}), "x.xlsx")


# Decision 125 (1): Aggre-SE wraps the sheet every 12 Cases; each band has its own No / 名称 and Khmax rows.
def test_the_cases_of_every_band_are_read_with_their_own_numbers_and_names():
    cases = {number: (f"{0.300 + 0.001 * number:.3f}", f"{0.400 + 0.001 * number:.3f}") for number in range(1, 16)}
    cases[14] = ("0.700", "0.699")
    found = read_max_response(workbook(cases), "x.xlsx")
    assert [(value.case, value.block) for value in found.values] == [
        *((number, "復旧性") for number in range(1, 16)), *((number, "安全性") for number in range(1, 16))]
    assert found.values[13] == ResponseValue(14, "線路方向(Case14_ρm10_αf1_正向)", "復旧性", found.values[13].kh)
    assert (found.adopted.case, found.adopted.block, found.kh.value) == (14, "復旧性", 0.700)
    assert found.values[-1].kh.value == 0.415  # Case 15 安全性: the third column of the second band


# Decision 125: the Cases of the 応答変位法 (from the stage's first_case on) are not the pier's response and are not read.
def test_the_cases_of_the_response_displacement_method_are_left_out_unread():
    cases = {number: ("0.431", "0.431") for number in range(1, 5)}
    cases |= {2: ("0.497", "0.497"), 4: ("0.497", "0.497")}
    cases |= {number: ("-----", "-----") for number in range(5, 16)}
    found = read_max_response(workbook(cases), "x.xlsx", response_from=5)
    assert [(value.case, value.block) for value in found.values] == [
        *((number, "復旧性") for number in range(1, 5)), *((number, "安全性") for number in range(1, 5))]
    assert (found.adopted.case, found.adopted.block, found.kh.value) == (2, "復旧性", 0.497)


def test_a_band_of_response_cases_alone_needs_no_khmax_rows():
    cases = {number: ("0.400", "0.401") for number in range(1, 13)} | {13: ("-----", "-----")}
    found = read_max_response(workbook(cases), "x.xlsx", response_from=13)
    assert len(found.values) == 24 and found.adopted.case == 1 and found.adopted.block == "安全性"


def test_a_used_case_that_is_not_a_number_still_stops_the_generation():
    cases = {1: ("0.3", "0.4"), 2: ("-----", "0.4"), 3: ("-----", "-----")}
    with pytest.raises(ResponseError, match=re.escape("x.xlsx Case2 復旧性 の最大応答震度 '-----' を数値として読めません")):
        read_max_response(workbook(cases), "x.xlsx", response_from=3)


def test_a_sheet_whose_cases_are_all_of_the_response_displacement_method():
    with pytest.raises(ResponseError, match=re.escape(f"x.xlsx の「{SHEET}」に Case がありません")):
        read_max_response(workbook({1: ("-----", "-----")}), "x.xlsx", response_from=1)


EXAMPLE4 = Path(__file__).parents[1] / "fixtures" / "example4" / "aggrese"


@pytest.mark.parametrize("file,first,khmax,adopted", [
    # Example橋脚4 post/線路方向_Aggre-SE.xlsx: 4 inertia Cases, then the 応答変位法 Cases 5〜11 (L1 '-----')
    ("線路方向_Aggre-SE.xlsx", 5, ("0.431", "0.497", "0.431", "0.497"), (2, "線路方向(Case2_ρm12_αf1_正向)", 0.497)),
    # post/線路直角方向_Aggre-SE.xlsx: 8 inertia Cases (正向・逆向), then Cases 9〜15 in two bands (13〜15 in the second)
    ("線路直角方向_Aggre-SE.xlsx", 9, ("0.695", "0.773", "0.695", "0.774", "0.695", "0.773", "0.695", "0.774"),
     (4, "線路直角方向(Case4_ρm12_αf2_正向)", 0.774)),
])
def test_the_maximum_response_of_example4_is_of_the_inertia_cases(file, first, khmax, adopted):
    found = read_max_response((EXAMPLE4 / file).read_bytes(), file, response_from=first)
    cases = range(1, first)
    assert [(value.case, value.block) for value in found.values] == [
        *((number, "復旧性") for number in cases), *((number, "安全性") for number in cases)]
    assert [str(value.kh.value) for value in found.values] == [*khmax, *khmax]  # 復旧性 = 安全性 in both workbooks
    assert (found.adopted.case, found.adopted.name, found.kh.value) == adopted and found.adopted.block == "復旧性"


@pytest.mark.parametrize("text", ["備考", "1.5"])
def test_a_no_row_with_a_cell_that_is_not_a_case_number(text):
    from io import BytesIO
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(workbook({1: ("0.3", "0.4")})))
    book[SHEET].cell(row=4, column=5, value=text)
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(ResponseError, match=re.escape(f"x.xlsx の「{SHEET}」の No の行に Case の番号でない値があります（{text}）")):
        read_max_response(stream.getvalue(), "x.xlsx")
