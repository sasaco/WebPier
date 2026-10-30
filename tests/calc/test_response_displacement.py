"""The 応答変位法の解析 of a pile foundation (decision 97, CONTEXT.md 応答変位法の解析).

The correction factors are fixed by Example橋脚4 (IV-56 表 3.6.3, IV-66〜69, IV-115〜118, IV-126〜129) and by hand:
α = Teq / (Tg / αg) to 3 decimals, νU = −2.0α + 2.5 and νL = −2.0α + 1.5 between 0.75 and 1.10, to 2 decimals.
"""

from decimal import Decimal
from io import BytesIO
import re

from openpyxl import Workbook
import pytest

from webpier.calc.response_displacement import (
    InertiaResponse, ResponseError, correction, read_inertia, response_cases, response_displacement)
from webpier.core.printed import Precision

S = Precision(4, "s")
TEQ = Precision(3, "s")
KH = Precision(3, "")


@pytest.mark.parametrize("level,teq,kh,kheq,alpha,upper,lower,adopted", [
    ("L1", 1.05, 0.378, 0.408, 1.378, 0.30, -0.70, 0.30),  # 線路方向 L1 (表 3.6.3)
    ("L2", 1.138, 0.428, 0.408, 1.067, 0.37, -0.63, 1.0),  # 線路方向 L2: past the bend, νU = 1.0 (解10.2.9)
    ("L1", 1.054, 0.378, 0.700, 1.383, 0.30, -0.70, 0.30),  # 直角方向 L1
    ("L2", 1.054, 0.706, 0.700, 0.988, 0.52, -0.48, 1.0),  # 直角方向 L2
    ("L1", 0.300, 0.2, 0.4, 0.394, 1.0, 0.0, 1.0),  # α ≦ 0.75
    ("L2", 0.400, 0.4, 0.4, 0.375, 1.0, 0.0, 1.0),  # the bend reached is not past it
])
def test_the_correction_factors_follow_the_sample(level, teq, kh, kheq, alpha, upper, lower, adopted):
    found = correction(level, TEQ(teq), S(0.5335), KH(kh), KH(kheq))
    assert (found.alpha.value, found.upper.value, found.lower.value, found.adopted.value) == (alpha, upper, lower, adopted)
    assert found.alpha_g.value == (0.7 if level == "L1" else 0.5)
    assert found.yielded is (kh > kheq)


def test_alpha_rounds_before_the_factors_and_the_factors_round_half_away():
    # 1.05 / (0.5 / 0.7) = 1.47 ≥ 1.10; 0.85 / (0.6 / 0.5) = 0.708333 → 0.708 ≤ 0.75; 1.0 / (1.0 / 0.925): 0.925 → 0.65, −0.35
    assert correction("L1", TEQ(1.05), S(0.5), KH(0.1), KH(0.2)).upper.value == 0.30
    assert correction("L2", TEQ(0.85), S(0.6), KH(0.1), KH(0.2)).alpha.value == 0.708
    found = correction("L2", TEQ(0.925), S(0.5), KH(0.1), KH(0.2))
    assert (found.alpha.value, found.upper.value, found.lower.value) == (0.925, 0.65, -0.35)


def test_the_cases_are_the_four_ways_and_one_of_two_alike_is_dropped():
    l1 = correction("L1", TEQ(1.05), S(0.5335), KH(0.378), KH(0.408))
    l2 = correction("L2", TEQ(1.138), S(0.5335), KH(0.428), KH(0.408))
    cases = response_cases(l1, l2)
    assert [(c.level, c.number, c.inertia.value, c.displacement.value, c.kh.value) for c in cases] == [
        ("L1", 1, 1.0, 0.30, 0.378), ("L1", 2, 1.0, -0.70, 0.378), ("L1", 3, 0.30, 1.0, 0.378),
        ("L1", 4, -0.70, 1.0, 0.378),
        ("L2", 1, 1.0, 1.0, 0.428), ("L2", 2, 1.0, -0.63, 0.428), ("L2", 4, -0.63, 1.0, 0.428)]  # 表 3.7.6 C6〜C8


def _workbook(columns, *, drop=None):
    """The summary of Aggre-SE with *columns* (name, αf, ρm, Khy, Kheq, Teq, 復旧性 Khmax, KhL1, 安全性 Khmax) for Cases 1, 2,
    ...; as Aggre-SE's, 12 Cases a band of 19 rows, each with its own headings (decision 125)."""
    book = Workbook()
    sheet = book.active
    sheet.title = "応答値一覧表"
    rows = [("地盤区分", ""), (None, None), ("応答値一覧表", None), ("解析ケース", "No"), (None, "名称"), (None, "αf"),
            (None, "ρm"), ("降伏点", "初期降伏震度"), (None, "全体系折曲点震度"), (None, "降伏変位"), (None, "降伏個所"),
            (None, "等価固有周期"), ("復旧性", "応答塑性率"), (None, "最大応答変位"), (None, "最大応答震度"),
            (None, "L1地震動"), ("安全性", "応答塑性率"), (None, "最大応答変位"), (None, "最大応答震度")]
    for band in range((len(columns) + 11) // 12 or 1):
        for number, (a, b) in enumerate(rows[2:] if band else rows, 3 if band else 1):
            if b == drop:
                continue
            sheet.cell(row=19 * band + number, column=1, value=a)
            sheet.cell(row=19 * band + number, column=2, value=b)
    for index, values in enumerate(columns):
        band, column = divmod(index, 12)
        for row, value in zip((4, 5, 6, 7, 8, 9, 12, 15, 16, 19), (index + 1, *values)):
            sheet.cell(row=19 * band + row, column=4 + column, value=value)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


CASE1 = ("P(Case1_ρm10_αf1_正向)", "1.0", "1.0", "0.386", "0.408", "0.982", "0.430", "0.321", "0.440")
CASE2 = ("P(Case2_ρm12_αf1_正向)", "1.0", "1.2", "0.449", "0.475", "0.994", "-----", "-----", "-----")
CASE3 = ("P_逆向(Case1_ρm10_αf1_逆向)", "1.0", "1.0", "0.386", "0.408", "0.982", "-----", "-----", "-----")


def test_the_inertia_response_is_read_from_the_case_of_the_summary():
    found = read_inertia(_workbook([CASE1, CASE2]), "線路方向_Aggre-SE.xlsx", 1)
    assert found == InertiaResponse(1, CASE1[0], TEQ(0.982), KH(0.408), KH(0.321), KH(0.440))


def test_a_case_of_the_second_band_is_read_with_its_own_rows():
    """Decision 125: Aggre-SE wraps the sheet every 12 Cases; a Case of the second band has its values there."""
    later = ("P(Case13_ρm10_αf1_正向)", "1.0", "1.0", "0.390", "0.410", "0.990", "0.431", "0.322", "0.441")
    found = read_inertia(_workbook([CASE2] * 12 + [later]), "線路方向_Aggre-SE.xlsx", 13)
    assert found == InertiaResponse(13, later[0], TEQ(0.990), KH(0.410), KH(0.322), KH(0.441))


@pytest.mark.parametrize("change,message", [
    (dict(case=3), "Case3 の「復旧性 最大応答震度」が数値ではありません（'-----'）"),
    (dict(case=4), "Case4 がありません"),
    (dict(case=2), "Case2 は ρm 1.0・αf 1.0 の Case ではありません"),
    (dict(drop="等価固有周期"), "「等価固有周期」の行がありません"),
    (dict(drop="L1地震動"), "「L1地震動」の行がありません"),
    (dict(alpha="2.0"), "Case1 は ρm 1.0・αf 1.0 の Case ではありません"),
    (dict(xlsx=b"not a zip"), "xlsx（zip）ではありません"),
])
def test_a_summary_without_the_values_stops_the_generation(change, message):
    case = change.get("case", 1)
    first = (*CASE1[:1], change.get("alpha", "1.0"), *CASE1[2:])
    data = change.get("xlsx") or _workbook([first, CASE2, CASE3], drop=change.get("drop"))
    with pytest.raises(ResponseError, match=re.escape(message)):
        read_inertia(data, "線路方向_Aggre-SE.xlsx", case)


def test_a_workbook_without_the_sheet_stops_the_generation():
    book = Workbook()
    book.active.title = "設計総括表"
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(ResponseError, match=re.escape("線路方向_Aggre-SE.xlsx にシート「応答値一覧表」がありません")):
        read_inertia(stream.getvalue(), "線路方向_Aggre-SE.xlsx", 1)


def test_the_result_takes_the_factors_of_both_levels_from_the_inertia_case():
    inertia = InertiaResponse(1, CASE1[0], TEQ(1.138), KH(0.408), KH(0.378), KH(0.428))
    result = response_displacement(inertia, S(0.5335), 9)
    assert [(c.level, c.alpha.value, c.adopted.value, c.lower.value) for c in result.corrections] == [
        ("L1", 1.493, 0.30, -0.70), ("L2", 1.067, 1.0, -0.63)]
    assert result.inertia is inertia and len(result.cases) == 7 and result.first_case == 9
