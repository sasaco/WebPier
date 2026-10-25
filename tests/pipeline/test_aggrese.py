"""The Aggre-SE runner of the 0711 / 0721 耐震性能照査 (decision 56, ADR-0009).

The runner is found from STRUCTURAL_MCP_AGGRESE_RUNNER (environment, then .env; the name structural-mcp uses) and
called as ``run`` with the project XML where it is staged and the DB1 of each Case mapped as the XML names them. Its
envelope, its five files and the Cases it aggregated are checked as the Aggre-SE runner's README states; the report
and the files published under post/ are returned as they are (decisions 56〜58, 60). Every case uses the stand-in of
tests/fixtures/aggrese/fake_runner.py (tests/conftest.py); the real runner is accepted by
scripts/accept_aggrese_report.py.
"""

from __future__ import annotations

import hashlib
import io
from pathlib import Path
import zipfile

import pymupdf
from openpyxl import load_workbook
import pytest

from webpier.pipeline import aggrese
from webpier.calc.footing import FoundationStep, StepsError
from webpier.pipeline.aggrese import (FOUNDATION_STEPS_COLUMNS, RUNNER_KEY, check_seismic, foundation_steps,
                                      runner_path, snap_map)
from webpier.pipeline.fempython import RunnerError
from webpier.reporting.stages.seismic_check import SEISMIC_CHECK

FOLDER = "線路方向_JRSNAP"
NAMES = ("線路方向(Case1_ρm10_αf1_正向).DB1", "線路方向(Case2_ρm10_αf1_正向).DB1")


def project(directory: Path, names=NAMES) -> Path:
    """A project XML as Aggre-SE writes it (Shift-JIS, declared utf-16) naming *names* in the JRSNAP folder, and the
    folder with a DB1 for each; the array ends with nil items as Aggre-SE's does."""
    (directory / FOLDER).mkdir(parents=True)
    for name in names:
        (directory / FOLDER / name).write_bytes(f"FAKE DB1 {name}".encode("utf-8"))
    items = "".join(f"\r\n      <string>{FOLDER}/{name}</string>" for name in names)
    text = ('<?xml version="1.0" encoding="utf-16"?>\r\n<Datas xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f"\r\n  <SNAP>\r\n    <SnapDBPath>{items}\r\n      <string xsi:nil=\"true\" />\r\n    </SnapDBPath>"
            "\r\n  </SNAP>\r\n</Datas>")
    path = directory / "線路方向.xml"
    path.write_bytes(text.encode("cp932"))
    return path


def texts(pdf: bytes) -> list[str]:
    with pymupdf.open(stream=pdf, filetype="pdf") as document:
        return [page.get_text().strip() for page in document]


def test_the_environment_comes_before_the_env_file(tmp_path: Path):
    runner, other = tmp_path / "runner.exe", tmp_path / "other.exe"
    runner.write_bytes(b""), other.write_bytes(b"")
    env = tmp_path / ".env"
    env.write_text(f"{RUNNER_KEY}='{other}'\nSTRUCTURAL_MCP_FEMPYTHON_RUNNER=x\n", encoding="utf-8")
    assert runner_path({RUNNER_KEY: str(runner)}, env) == runner
    assert runner_path({RUNNER_KEY: ""}, env) == other
    assert runner_path({}, env) == other


def test_the_headless_runner_reads_only_the_environment(tmp_path: Path):
    """Decision 61: without dotenv a runner set only in the .env file is not set."""
    runner, other = tmp_path / "runner.exe", tmp_path / "other.exe"
    runner.write_bytes(b""), other.write_bytes(b"")
    env = tmp_path / ".env"
    env.write_text(f"{RUNNER_KEY}='{other}'\n", encoding="utf-8")
    assert runner_path({RUNNER_KEY: str(runner)}, env, dotenv=False) == runner
    with pytest.raises(RunnerError, match=f"環境変数 {RUNNER_KEY} に AggreSE.Headless.exe の絶対パスを設定してください"):
        runner_path({}, env, dotenv=False)


@pytest.mark.parametrize("setting,message", [(None, "Aggre-SE の runner が設定されていません"), ("", "設定されていません"),
                                             ("runner.exe", "絶対パス"), ("missing", "ありません")])
def test_a_runner_that_is_not_set_up_is_an_error(tmp_path: Path, setting, message):
    env = tmp_path / ".env"
    if setting is not None:
        value = str(tmp_path / "missing.exe") if setting == "missing" else setting
        env.write_text(f"{RUNNER_KEY}={value}\n", encoding="utf-8")
    with pytest.raises(RunnerError, match=message):
        runner_path({}, env)


def test_the_process_environment_is_used_by_default(monkeypatch: pytest.MonkeyPatch, fake_aggrese_runner: Path):
    monkeypatch.setenv(RUNNER_KEY, str(fake_aggrese_runner))
    assert runner_path() == fake_aggrese_runner


def test_the_snap_map_is_each_db1_the_xml_names_from_its_folder(tmp_path: Path):
    path = project(tmp_path)
    assert snap_map(path) == {str(index): str((tmp_path / FOLDER / name).resolve())
                              for index, name in enumerate(NAMES)}


@pytest.mark.parametrize("content,message", [
    (b'<?xml version="1.0" encoding="utf-16"?>\r\n<Datas><SNAP><SnapDBPath /></SNAP></Datas>', "SnapDBPath"),
    (b'<?xml version="1.0" encoding="utf-16"?>\r\n<Datas />', "SnapDBPath"),
    (b"<Datas>", "プロジェクトを読めません"), (b"\x82", "プロジェクトを読めません"),
])
def test_a_project_without_databases_is_an_error(tmp_path: Path, fake_aggrese_runner: Path, content, message):
    path = tmp_path / "線路方向.xml"
    path.write_bytes(content)
    with pytest.raises(RunnerError, match=f"線路方向.xml で耐震性能照査ができません: .*{message}"):
        check_seismic(path, fake_aggrese_runner)


def test_the_stand_in_checks_each_case_the_project_names(tmp_path: Path, fake_aggrese_runner: Path):
    path = project(tmp_path)
    checked = check_seismic(path, fake_aggrese_runner)
    assert checked.engine_version == "fake-1"
    assert [(case["caseNumber"], case["name"], case["aggregated"]) for case in checked.cases] == [
        (1, "線路方向(Case1_ρm10_αf1_正向)", True), (2, "線路方向(Case2_ρm10_αf1_正向)", True)]
    pages = texts(checked.pdf)
    assert pages == [*(f"FAKE AGGRESE Case{number} "
                       f"{hashlib.sha256((tmp_path / FOLDER / name).read_bytes()).hexdigest()[:16]}"
                       for number, name in enumerate(NAMES, 1)),
                     f"FAKE AGGRESE SUMMARY {hashlib.sha256(path.read_bytes()).hexdigest()[:16]}"]
    assert checked.pages == {SEISMIC_CHECK: checked.pdf}
    digest = hashlib.sha256(path.read_bytes()).hexdigest()[:16]
    assert sorted(checked.files) == ["pik", "steps", "xlsx"]
    assert checked.files["pik"] == f"Batch ver1.0.0\r\nFAKE PIK {digest}\r\n".encode("cp932")
    # Decision 59: a real workbook whose 応答値一覧表 holds the 最大応答震度 of each Case (0.385/0.386, then +0.01 a Case).
    workbook = load_workbook(io.BytesIO(checked.files["xlsx"]), read_only=True)
    assert workbook.sheetnames == ["応答値一覧表", "設計総括表", "安全率一覧表", "直接基礎"]
    sheet = workbook["応答値一覧表"]
    assert (sheet["B15"].value, sheet["D15"].value, sheet["E15"].value) == ("最大応答震度", "0.385", "0.395")
    assert (sheet["A17"].value, sheet["D19"].value, sheet["E19"].value) == ("安全性", "0.386", "0.396")
    assert workbook["設計総括表"]["A1"].value == f"FAKE XLSX {digest}"
    workbook.close()
    # Decision 60: the direct-foundation step table, UTF-8 and CRLF, one row per Case, node and output step.
    steps = checked.files["steps"].decode("utf-8")
    assert "\n" not in steps.replace("\r\n", "") and steps.endswith("\r\n")
    lines = steps.split("\r\n")[:-1]
    assert lines[0] == ",".join(FOUNDATION_STEPS_COLUMNS) == (
        "case_index,case_name,node,step,V,H,M,qvd,qd,Rvd,Mmd,plasticity")
    assert len(lines) == 1 + 4 * len(NAMES)
    assert lines[1] == "0,線路方向(Case1_ρm10_αf1_正向),7,1,16825.97,0,580.25,8650.41,8710.57,543358.71,57085.22,0"
    assert lines[7] == "1,線路方向(Case2_ρm10_αf1_正向),7,3,16825.97,-5901.07,54268.43,3236.74,3296.9,16655.87,54120.2,0.081"
    assert sorted(path.name for path in (tmp_path / FOLDER).iterdir()) == sorted(NAMES)


@pytest.mark.parametrize("mode,message", [
    ("fail", "CALCULATION_FAILED: 集計できませんでした。"),
    ("exit", "JSON ではありません"), ("text", "JSON ではありません"), ("list", "オブジェクトではありません"),
    ("protocol", "やり取りの版が違います"), ("engine", "やり取りの版が違います"), ("version", "engineVersion"),
    ("artifacts", "成果物が結果・計算書・ピックアップ・総括表・直接基礎のステップ表の5つではありません"),
    ("nopik", "成果物が結果・計算書・ピックアップ・総括表・直接基礎のステップ表の5つではありません"),
    ("noxlsx", "成果物が結果・計算書・ピックアップ・総括表・直接基礎のステップ表の5つではありません"),
    ("notzip", r"総括表が xlsx（zip）ではありません"),
    ("nocsv", "成果物が結果・計算書・ピックアップ・総括表・直接基礎のステップ表の5つではありません"),
    ("badcsv", "直接基礎のステップ表の列が違います"), ("notutf8csv", "直接基礎のステップ表を UTF-8 で読めません"),
    ("missing", "成果物 report.pdf がありません"), ("extra", "出力先のファイルが成果物の一覧と一致しません"),
    ("sha", r"report\.pdf が runner の応答の SHA256・バイト数と一致しません"),
    ("cases", r"解析ケース（summary\.cases）がありません"), ("unaggregated", "全 2 ケースが集計されていません"),
    ("broken", "計算書の PDF を読めません"), ("empty", "計算書の PDF にページがありません"),
])
def test_a_runner_that_does_not_check_is_an_error(tmp_path: Path, fake_aggrese_runner: Path,
                                                  monkeypatch: pytest.MonkeyPatch, mode, message):
    path = project(tmp_path)
    monkeypatch.setenv("FAKE_AGGRESE", mode)
    with pytest.raises(RunnerError, match=f"線路方向.xml で耐震性能照査ができません: .*{message}"):
        check_seismic(path, fake_aggrese_runner)


def test_a_missing_database_is_the_runners_error(tmp_path: Path, fake_aggrese_runner: Path):
    path = project(tmp_path)
    (tmp_path / FOLDER / NAMES[1]).unlink()
    with pytest.raises(RunnerError, match="SNAP_DB_NOT_FOUND"):
        check_seismic(path, fake_aggrese_runner)


def test_a_runner_that_cannot_start_is_an_error(tmp_path: Path):
    path = project(tmp_path)
    runner = tmp_path / "runner.exe"
    runner.write_bytes(b"not a program")
    with pytest.raises(RunnerError, match="runner を起動できません"):
        check_seismic(path, runner)


def test_a_runner_that_does_not_finish_is_an_error(tmp_path: Path, fake_aggrese_runner: Path,
                                                   monkeypatch: pytest.MonkeyPatch):
    path = project(tmp_path)
    monkeypatch.setattr(aggrese.subprocess, "run", _timeout)
    with pytest.raises(RunnerError, match="秒で終わりませんでした"):
        check_seismic(path, fake_aggrese_runner, timeout=1)


def _timeout(*args, **kwargs):
    raise aggrese.subprocess.TimeoutExpired(args[0], kwargs.get("timeout", 1))


HEADER = "case_index,case_name,node,step,V,H,M,qvd,qd,Rvd,Mmd,plasticity"
ROW = "0,線路方向(Case1_ρm10_αf1_正向),7,32,16825.970703125,-5901.0673828125,54168.42578125,3236.7393900986663," \
      "3296.8993900986661,16655.873658566576,54120.199464658981,0.081009116849374715"


def test_the_step_table_is_read_into_rows(tmp_path: Path, fake_aggrese_runner: Path):
    """Decision 60: the step table of the direct foundation, as the runner writes it (UTF-8, CRLF), row by row."""
    rows = foundation_steps(f"{HEADER}\r\n{ROW}\r\n1,Case2,8,1,1.5,-2,3,4,5,6,7,0\r\n".encode("utf-8"))
    assert rows == (
        FoundationStep(0, "線路方向(Case1_ρm10_αf1_正向)", 7, 32, 16825.970703125, -5901.0673828125, 54168.42578125,
                       3236.7393900986663, 3296.8993900986661, 16655.873658566576, 54120.199464658981,
                       0.081009116849374715),
        FoundationStep(1, "Case2", 8, 1, 1.5, -2.0, 3.0, 4.0, 5.0, 6.0, 7.0, 0.0))
    checked = check_seismic(project(tmp_path), fake_aggrese_runner)
    stand_in = foundation_steps(checked.files["steps"])
    assert [(row.case_index, row.node, row.step) for row in stand_in] == [(index, 7, step) for index in (0, 1)
                                                                          for step in (1, 2, 3, 4)]
    assert stand_in[2].v >= stand_in[2].rvd > stand_in[1].rvd - stand_in[1].v  # step 3 reaches Rvd


def test_a_bom_and_lf_endings_are_tolerated():
    (row,) = foundation_steps(("﻿" + HEADER + "\n" + ROW + "\n").encode("utf-8"))
    assert (row.case_index, row.step, row.v) == (0, 32, 16825.970703125)


@pytest.mark.parametrize("content,message", [
    (b"case,step\r\n0,1\r\n", "直接基礎のステップ表の列が違います"),
    (f"{HEADER}\r\n".encode(), "直接基礎のステップ表に行がありません"),
    (f"{HEADER}\r\n0,Case1,7,x,1,2,3,4,5,6,7,0\r\n".encode(), "直接基礎のステップ表を読めません"),
    (f"{HEADER}\r\n0,Case1,7,1,1,2,3,4,5,6,7\r\n".encode(), "直接基礎のステップ表を読めません"),
    (f"{HEADER}\r\n0,Case1,7,1,nan,2,3,4,5,6,7,0\r\n".encode(), "有限でない値"),
    ("直接基礎\r\n".encode("cp932"), "直接基礎のステップ表を読めません"),
])
def test_an_unreadable_step_table_is_an_error(content, message):
    with pytest.raises(StepsError, match=message):
        foundation_steps(content)


# --- a pile foundation (decision 94): the limit values of the stability are the .sdr's, and there is no step table ----------


def pile_project(directory: Path, foundation: int = 0) -> Path:
    """The project of *project* with a ``FoundationType`` (0 piles, 1 direct), as Aggre-SE writes it."""
    path = project(directory)
    text = path.read_bytes().decode("cp932").replace(
        "</Datas>", f"  <Foundation>\r\n    <基礎照査>\r\n      <clsFoundation>\r\n        <FoundationType>{foundation}"
                    "</FoundationType>\r\n      </clsFoundation>\r\n    </基礎照査>\r\n  </Foundation>\r\n</Datas>")
    path.write_bytes(text.encode("cp932"))
    return path


SDR = b"Ver.5.2.3\r\nWebPier\r\n"


def test_a_pile_foundation_is_checked_with_the_sdr_and_returns_no_step_table(tmp_path: Path, fake_aggrese_runner: Path):
    path = pile_project(tmp_path)
    checked = check_seismic(path, fake_aggrese_runner, limits=aggrese.Limits(SDR, aggrese.AXIAL))
    assert sorted(checked.files) == ["pik", "xlsx"]
    assert texts(checked.pdf)[-1] == f"FAKE AGGRESE LIMITS axial {hashlib.sha256(SDR).hexdigest()[:16]}"
    assert checked.pages == {SEISMIC_CHECK: checked.pdf}
    assert aggrese.PILE_PUBLISHED == {kind: item for kind, item in aggrese.PUBLISHED.items() if kind != "steps"}
    other = check_seismic(path, fake_aggrese_runner, limits=aggrese.Limits(SDR, aggrese.TRANSVERSE))
    assert texts(other.pdf)[-1].startswith("FAKE AGGRESE LIMITS transverse ")


def test_the_direct_foundation_has_no_limits_and_a_step_table(tmp_path: Path, fake_aggrese_runner: Path):
    checked = check_seismic(pile_project(tmp_path, foundation=1), fake_aggrese_runner)
    assert sorted(checked.files) == ["pik", "steps", "xlsx"]
    assert not any("LIMITS" in text for text in texts(checked.pdf))


def test_a_pile_foundation_without_the_sdr_and_a_direct_one_with_it_are_refused(tmp_path: Path, fake_aggrese_runner: Path):
    with pytest.raises(RunnerError, match="SDR_REQUIRED"):
        check_seismic(pile_project(tmp_path / "a"), fake_aggrese_runner)
    with pytest.raises(RunnerError, match="INVALID_ARGUMENT"):
        check_seismic(pile_project(tmp_path / "b", foundation=1), fake_aggrese_runner,
                      limits=aggrese.Limits(SDR, aggrese.AXIAL))


@pytest.mark.parametrize("mode,message", [("nopik", "結果・計算書・ピックアップ・総括表の4つ"), ("noxlsx", "4つ"),
                                          ("missing", "report.pdf がありません"), ("sha", "SHA256")])
def test_a_pile_foundation_runner_that_misbehaves_is_an_error(tmp_path: Path, fake_aggrese_runner: Path,
                                                              monkeypatch: pytest.MonkeyPatch, mode, message):
    monkeypatch.setenv("FAKE_AGGRESE", mode)
    with pytest.raises(RunnerError, match=message):
        check_seismic(pile_project(tmp_path), fake_aggrese_runner, limits=aggrese.Limits(SDR, aggrese.AXIAL))
