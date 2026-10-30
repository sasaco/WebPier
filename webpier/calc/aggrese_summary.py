"""The sheet 応答値一覧表 of Aggre-SE's summary workbook, band by band (decisions 59, 97, 125).

Aggre-SE writes 12 Cases a band and wraps the rest into further bands below (Example橋脚4 線路直角方向: Cases 1〜12, then
13〜15 from row 22). Each band starts with its own row No (column B) of the Case numbers from column D and has its own
名称 and value rows; column A labels the block (解析ケース, 降伏点, 復旧性, 安全性) of the rows from it on. A row is found by
its block and heading, never by its number, and a Case's values are in its band's rows at its column.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import zipfile

from openpyxl import load_workbook

SHEET = "応答値一覧表"
FIRST_CASE_COLUMN = 3  # column D, as an index from column A


@dataclass(frozen=True)
class Band:
    """One band of the sheet: its Cases and its rows as text."""

    cases: tuple[tuple[int, int], ...]  # (column index, Case number), from column D
    lines: tuple[tuple[str, tuple[str, ...]], ...]  # (block, the row's cell texts, as wide as the sheet)

    def row(self, heading: str, block: str | None = None) -> tuple[str, ...] | None:
        """The first row of the band headed *heading* in column B (in the block *block*, or in any)."""
        return next((cells for label, cells in self.lines if cells[1] == heading and block in (None, label)), None)


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def summary_bands(xlsx: bytes, file: str, error: type[ValueError]) -> tuple[Band, ...]:
    """The bands of the sheet of the summary workbook bytes *xlsx* (published as *file*); *error* when the bytes are not
    a workbook with the sheet. The rows above the first No row belong to no band."""
    if not zipfile.is_zipfile(BytesIO(xlsx)):
        raise error(f"{file} が xlsx（zip）ではありません")
    try:
        workbook = load_workbook(BytesIO(xlsx), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several kinds for a damaged workbook
        raise error(f"{file} を読めません（{exc}）") from exc
    if SHEET not in workbook.sheetnames:
        raise error(f"{file} にシート「{SHEET}」がありません")
    sheet = workbook[SHEET]
    width = max(sheet.max_column, 2)
    rows = [tuple(_text(value) for value in row) + ("",) * (width - len(row))
            for row in sheet.iter_rows(min_col=1, max_col=width, values_only=True)]
    workbook.close()
    found: list[tuple[tuple[str, ...], list[tuple[str, tuple[str, ...]]]]] = []
    block = ""
    for cells in rows:
        block = cells[0] or block
        if cells[1] == "No":
            found.append((cells, []))
        if found:
            found[-1][1].append((block, cells))
    return tuple(Band(tuple((index, int(float(text))) for index, text in enumerate(numbers)
                            if index >= FIRST_CASE_COLUMN and text), tuple(lines))
                 for numbers, lines in found)


__all__ = ["Band", "SHEET", "summary_bands"]
